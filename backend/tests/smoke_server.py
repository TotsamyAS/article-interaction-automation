"""Run inside a disposable Docker container; never against experiment storage."""
import csv
import io
import json
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
import uuid
import zipfile
from pathlib import Path

from openpyxl import load_workbook


def main():
    with tempfile.TemporaryDirectory(prefix="interaction-smoke-") as directory:
        bootstrap = Path(directory) / "server.py"
        bootstrap.write_text(
            "import sys\n"
            "sys.path.insert(0, '/app')\n"
            "import uvicorn\n"
            "from app.api import create_app\n"
            "from app.config import load_settings\n"
            "from app.access import Principal, Signatures, require_principal\n"
            "from unittest.mock import Mock\n"
            f"settings = load_settings().model_copy(update={{'database_path': {str(Path(directory) / 'smoke.sqlite3')!r}}})\n"
            "app = create_app(settings, signatures_factory=lambda: Mock(spec=Signatures))\n"
            "app.dependency_overrides[require_principal] = lambda: Principal('smoke', 'SMOKE', 'researcher', 1)\n"
            "uvicorn.run(app, host='127.0.0.1', port=18081, access_log=False, log_level='error')\n",
            encoding="utf-8",
        )
        process = subprocess.Popen([sys.executable, str(bootstrap)])
        base = "http://127.0.0.1:18081"

        def request(path, body=None):
            payload = json.dumps(body).encode() if body is not None else None
            req = urllib.request.Request(base + path, data=payload,
                                         headers={"Content-Type": "application/json"})
            with urllib.request.urlopen(req, timeout=10) as response:
                return response.read()

        def data(path, body=None):
            return json.loads(request(path, body))

        try:
            for _ in range(100):
                if process.poll() is not None:
                    raise RuntimeError("Server exited before becoming ready")
                try:
                    assert data("/health")["status"] == "ok"
                    break
                except urllib.error.URLError:
                    time.sleep(0.1)
            else:
                raise RuntimeError("Server readiness timed out")
            session = data("/api/sessions", {"participant_code": "SMOKE", "kind": "experiment"})
            trial_id = session["trials"][0]["id"]
            data(f"/api/trials/{trial_id}/start", {})
            query = {"filters": [{"field": "status", "operator": "eq", "value": "В работе"},
                                  {"field": "priority", "operator": "eq", "value": "Высокий"}]}
            response = data(f"/api/trials/{trial_id}/attempts", {"request_id": str(uuid.uuid4()), "query": query})
            assert response["correct"]
            rows = list(csv.DictReader(io.StringIO(request("/api/analytics/trials.csv?completed_only=true").decode("utf-8-sig"))))
            assert len(rows) == 1 and rows[0]["status"] == "correct"
            workbook = load_workbook(io.BytesIO(request("/api/analytics/export.xlsx")), read_only=True)
            assert set(workbook.sheetnames) == {"protocol", "trials", "attempts", "events", "summary"}
            workbook.close()
            with zipfile.ZipFile(io.BytesIO(request("/api/analytics/export.zip"))) as archive:
                assert "protocol.json" in archive.namelist()
            print("PASS: real HTTP server, migrations, successful attempt, analytics CSV/XLSX/ZIP")
        finally:
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()


if __name__ == "__main__":
    main()
