import csv
import io
import json
import zipfile
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.analytics import excel_cell
from app.errors import DomainError


def create_started(client, code="P001", kind="experiment"):
    response = client.post("/api/sessions", json={"participant_code": code, "kind": kind})
    assert response.status_code == 200
    session = response.json()
    trial = session["trials"][0]
    assert client.post(f'/api/trials/{trial["id"]}/start').status_code == 200
    return session, trial


def test_api_opens_and_exposes_typed_contracts(client):
    assert client.get("/health").json() == {"status": "ok"}
    schema = client.get("/openapi.json")
    assert schema.status_code == 200
    assert "Query" in schema.json()["components"]["schemas"]
    assert len(client.get("/api/records").json()) == 120
    assert len(client.get("/api/tasks").json()) == 16


def test_analytics_exports_preserve_actual_and_penalty_and_exclude_training(client, clock):
    _, trial = create_started(client)
    response = client.post(f'/api/trials/{trial["id"]}/attempts', json={"request_id": str(uuid4()), "query": {}})
    assert response.status_code == 200
    clock.advance(1500000)
    create_started(client, "TRAINING", "practice")
    snapshot = client.get("/api/analytics?completed_only=true").json()
    assert len(snapshot["trials"]) == 1
    row = snapshot["trials"][0]
    assert row["Nretry_actual"] == 0 and row["Nretry_analysis"] == 24
    assert row["trial_limit_seconds"] == 1500 and row["attempt_limit"] == 25
    assert row["Tcorrect_actual_ms"] is None and row["Tcorrect_analysis_ms"] == 1500000
    assert snapshot["summary"][0]["Tcorrect_actual_mean_ms"] is None
    assert snapshot["summary"][0]["Tcorrect_analysis_mean_ms"] == 1500000
    assert snapshot["protocol"]["training_excluded"]

    raw = client.get("/api/analytics/trials.csv?completed_only=true")
    assert raw.status_code == 200
    rows = list(csv.DictReader(io.StringIO(raw.content.decode("utf-8-sig"))))
    assert rows[0]["Nretry_actual"] == "0" and rows[0]["Nretry_analysis"] == "24"

    archive_response = client.get("/api/analytics/export.zip?completed_only=true")
    assert archive_response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(archive_response.content)) as archive:
        assert set(archive.namelist()) == {"registry.csv", "trials.csv", "attempts.csv", "events.csv", "interpretations.csv", "m4_transcriptions.csv", "agent_runs.csv", "summary.csv", "protocol.json"}
        assert json.loads(archive.read("protocol.json"))["manifest"]["dataset_sha256"]

    workbook_response = client.get("/api/analytics/export.xlsx?completed_only=true")
    assert workbook_response.status_code == 200
    workbook = load_workbook(io.BytesIO(workbook_response.content), read_only=True)
    assert set(workbook.sheetnames) == {"protocol", "registry", "trials", "attempts", "events", "interpretations", "m4_transcriptions", "agent_runs", "summary"}
    values = list(workbook["trials"].values)
    excel_row = dict(zip(values[0], values[1]))
    assert excel_row["Nretry_actual"] == 0
    assert excel_row["Nretry_analysis"] == 24
    assert excel_row["Tcorrect_actual_ms"] is None
    workbook.close()


def test_analytics_filters_and_empty_exports(client):
    create_started(client)
    create_started(client, "T001", "practice")
    response = client.get("/api/analytics?participant_code=P001&mode=M1&level=1")
    assert response.status_code == 200, response.json()
    snapshot = response.json()
    assert len(snapshot["trials"]) == 1
    assert snapshot["trials"][0]["participant_code"] == "P001"
    assert snapshot["summary"][0]["completed"] == 0
    assert snapshot["summary"][0]["success_rate"] is None
    all_rows = client.get("/api/analytics?include_practice=true").json()["trials"]
    assert len(all_rows) == 20
    empty = client.get("/api/analytics/trials.csv?participant_code=ABSENT")
    assert empty.status_code == 200
    assert len(list(csv.reader(io.StringIO(empty.content.decode("utf-8-sig"))))) == 1


def test_successful_api_flow_idempotency_and_session_export(client, clock):
    session, trial = create_started(client)
    query = {"filters": [{"field": "status", "operator": "eq", "value": "В работе"},
                          {"field": "priority", "operator": "eq", "value": "Высокий"}]}
    preview = client.post(f'/api/trials/{trial["id"]}/preview', json=query)
    assert preview.status_code == 200
    clock.advance(1000)
    payload = {"request_id": str(uuid4()), "query": query}
    first = client.post(f'/api/trials/{trial["id"]}/attempts', json=payload)
    repeated = client.post(f'/api/trials/{trial["id"]}/attempts', json=payload)
    assert first.status_code == 200 and first.json()["correct"]
    assert first.json() == repeated.json()
    saved = client.get(f'/api/sessions/{session["id"]}/export')
    assert len(saved.json()["attempt_log"]) == 1
    assert saved.json()["trials"][0]["metrics"]["actual"]["Tcorrect_ms"] == 1000


def test_task_mining_events_export_explicit_participant_order_and_durations(client, clock):
    _, trial = create_started(client, code="P-DURATION")
    started = client.get(f'/api/trials/{trial["id"]}').json()["started_ms"]
    request_id = str(uuid4())
    clock.advance(10000)
    raw_events = [
        {"sequence": 0, "offset_ms": 1000, "kind": "input", "target": "m2-builder", "action": "click"},
        {"sequence": 1, "offset_ms": 2000, "kind": "focus_lost", "target": "window"},
        {"sequence": 2, "offset_ms": 3500, "kind": "focus_gained", "target": "window"},
        {"sequence": 3, "offset_ms": 4000, "kind": "request_started", "target": "m2-preview", "request_id": request_id},
        {"sequence": 4, "offset_ms": 6000, "kind": "request_finished", "target": "m2-preview", "request_id": request_id},
        {"sequence": 5, "offset_ms": 6500, "kind": "speech_started", "target": "m4-speech"},
        {"sequence": 6, "offset_ms": 8000, "kind": "speech_finished", "target": "m4-speech"},
        {"sequence": 7, "offset_ms": 8500, "kind": "navigation", "target": "m2-mode-hint:open"},
        {"sequence": 8, "offset_ms": 9000, "kind": "navigation", "target": "m2-mode-hint:closed"},
        {"sequence": 9, "offset_ms": 9500, "kind": "navigation", "target": "m2-mode-hint:open"},
    ]
    events = [{"event_id": str(uuid4()), **event} for event in raw_events]
    response = client.post(f'/api/trials/{trial["id"]}/events', json={"events": events})
    assert response.status_code == 200, response.json()

    snapshot = client.get('/api/analytics?participant_code=P-DURATION').json()
    rows = snapshot["events"]
    assert len(rows) == len(events)
    assert {row["participant_code"] for row in rows} == {"P-DURATION"}
    assert [row["sequence"] for row in rows] == list(range(len(events)))
    assert [row["occurred_ms"] for row in rows] == [started + event["offset_ms"] for event in raw_events]
    assert rows[0]["duration_ms"] == 0 and rows[0]["duration_complete"] is True
    assert rows[1]["duration_ms"] == 1500 and rows[1]["duration_complete"] is True
    assert rows[3]["duration_ms"] == 2000 and rows[3]["duration_complete"] is True
    assert rows[5]["duration_ms"] == 1500 and rows[5]["duration_complete"] is True
    assert rows[7]["duration_ms"] == 500 and rows[7]["duration_complete"] is True
    assert rows[9]["duration_ms"] == 500 and rows[9]["duration_complete"] is False
    assert all(row["duration_ms"] >= 0 for row in rows)
    assert snapshot["protocol"]["export_schema_version"] == 8

    csv_response = client.get('/api/analytics/events.csv?participant_code=P-DURATION')
    exported = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert exported[0]["participant_code"] == "P-DURATION"
    assert exported[1]["duration_ms"] == "1500"
    assert exported[1]["occurred_ms"] == str(started + 2000)


def test_validation_does_not_echo_untrusted_request(client):
    response = client.post("/api/sessions", json={"participant_code": "<invalid-value>"})
    assert response.status_code == 422
    assert "<invalid-value>" not in response.text


@pytest.mark.parametrize("level", ["0", "4", "invalid"])
def test_analytics_rejects_invalid_level(client, level):
    assert client.get(f"/api/analytics?level={level}").status_code == 422


def test_excel_never_silently_truncates_or_executes_text():
    assert excel_cell("=1+1") == "'=1+1"
    assert excel_cell("x\x00y") == "x\\u0000y"
    with pytest.raises(DomainError, match="лимит"):
        excel_cell("x" * 32768)
