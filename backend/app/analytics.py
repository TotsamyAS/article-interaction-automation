import csv
import io
import zipfile
from collections import defaultdict
from statistics import mean

from openpyxl import Workbook
from pydantic import Field

from .contracts import Contract, Mode
from .database import encode
from .errors import DomainError
from .service import ExperimentService


class AnalyticsFilter(Contract):
    participant_code: str | None = None
    mode: Mode | None = None
    level: int | None = Field(default=None, ge=1, le=3)
    include_practice: bool = False
    completed_only: bool = False


TABLE_COLUMNS = {
    "registry": ["participant_code", "session_id", "sequence_no", "created_ms", "complete", "completion_code", "trials", "correct", "incomplete"],
    "trials": ["session_id", "participant_code", "kind", "sequence_no", "trial_id", "position", "block_index", "mode", "task_id", "level", "tci",
               "status", "end_reason", "gave_up", "trial_limit_seconds", "attempt_limit", "wording_version", "started_ms", "ended_ms", "elapsed_ms", "Tcorrect_actual_ms", "Tfirst_ms", "Tuser_active_ms",
               "A1_actual", "attempts", "Nretry_actual", "Tcorrect_analysis_ms", "A1_analysis", "Nretry_analysis", "incomplete"],
    "attempts": ["session_id", "participant_code", "kind", "trial_id", "mode", "task_id", "level", "attempt_id", "ordinal", "correct",
                 "started_ms", "finished_ms", "Texec_ms", "query", "record_ids", "aggregate", "selected_groups", "export_url"],
    "events": ["session_id", "participant_code", "kind", "trial_id", "trial_position", "mode", "task_id", "level", "event_id", "sequence", "offset_ms",
               "occurred_ms", "duration_ms", "duration_complete", "received_ms", "event_kind", "target", "action", "x", "y", "key", "request_id"],
    "interpretations": ["session_id", "participant_code", "kind", "trial_id", "mode", "task_id", "level", "request_id", "user_text",
                        "status", "requested_model", "response_model", "provider", "system_fingerprint", "prompt_version", "prompt_sha256", "temperature",
                        "raw_response", "query", "error_code", "error_message",
                        "Tllm_ms", "input_tokens", "output_tokens", "started_ms", "finished_ms"],
    "m4_transcriptions": ["session_id", "participant_code", "kind", "trial_id", "mode", "task_id", "level", "request_id",
                          "mime_type", "audio_bytes", "audio_duration_ms", "requested_model", "compute_type", "requested_language",
                          "detected_language", "language_probability", "transcript", "Tasr_ms", "started_ms", "finished_ms"],
    "agent_runs": ["session_id", "participant_code", "kind", "trial_id", "mode", "task_id", "level", "request_id", "user_text",
                   "status", "requested_model", "response_model", "provider", "prompt_version", "prompt_sha256", "temperature", "max_steps",
                   "query", "trajectory", "final_text", "termination", "error_code", "error_message",
                   "Tllm_ms", "Ttool_ms", "Tagent_ms", "llm_calls", "tool_calls", "input_tokens", "output_tokens", "started_ms", "finished_ms"],
    "summary": ["kind", "mode", "level", "trial_limit_seconds", "attempt_limit", "wording_version", "participants", "trials", "pending", "active", "completed", "correct", "incomplete", "success_rate",
                "A1_observed_n", "A1_actual_rate", "A1_analysis_rate", "Tcorrect_actual_n", "Tcorrect_actual_mean_ms",
                "Tcorrect_analysis_mean_ms", "Tuser_active_mean_ms", "Nretry_actual_mean", "Nretry_analysis_mean"],
}


def _event_span(payload: dict) -> tuple[tuple[str, str], bool] | None:
    kind = payload.get("kind")
    if kind in ("focus_lost", "focus_gained"):
        return (("focus", "window"), kind == "focus_lost")
    if kind in ("request_started", "request_finished") and payload.get("request_id"):
        return (("request", str(payload["request_id"])), kind == "request_started")
    if kind in ("speech_started", "speech_finished"):
        return (("speech", payload.get("target", "")), kind == "speech_started")
    if kind == "navigation":
        target = payload.get("target", "")
        if target.endswith(":open"):
            return (("navigation", target[:-5]), True)
        if target.endswith(":closed"):
            return (("navigation", target[:-7]), False)
    return None


def event_duration_metadata(event_log: list[dict], trials_by_id: dict[str, dict]) -> dict[tuple[str, str], dict]:
    """Calculate event durations while keeping censored intervals distinguishable."""
    metadata: dict[tuple[str, str], dict] = {}
    open_spans: dict[tuple[str, tuple[str, str]], tuple[tuple[str, str], int]] = {}
    for event in event_log:
        event_key = (event["trial_id"], event["event_id"])
        metadata[event_key] = {"duration_ms": 0, "duration_complete": True}
        span = _event_span(event["payload"])
        if span is None:
            continue
        span_key, starts = span
        key = (event["trial_id"], span_key)
        if starts:
            previous = open_spans.get(key)
            if previous is not None:
                previous_key, previous_offset = previous
                metadata[previous_key] = {
                    "duration_ms": max(0, event["offset_ms"] - previous_offset),
                    "duration_complete": False,
                }
            open_spans[key] = (event_key, event["offset_ms"])
            metadata[event_key] = {"duration_ms": 0, "duration_complete": False}
            continue
        previous = open_spans.pop(key, None)
        if previous is not None:
            previous_key, previous_offset = previous
            metadata[previous_key] = {
                "duration_ms": max(0, event["offset_ms"] - previous_offset),
                "duration_complete": True,
            }

    for (trial_id, _), (event_key, start_offset) in open_spans.items():
        observation_end = max(start_offset, trials_by_id[trial_id]["elapsed_since_start_ms"])
        metadata[event_key] = {
            "duration_ms": observation_end - start_offset,
            "duration_complete": False,
        }
    return metadata


class AnalyticsService:
    def __init__(self, experiment: ExperimentService):
        self.experiment = experiment

    def collect(self, filters: AnalyticsFilter) -> dict:
        registry, trials, attempts, events, interpretations, m4_transcriptions, agent_runs = [], [], [], [], [], [], []
        # One transaction freezes a consistent snapshot of all exported tables.
        with self.experiment.database.transaction() as connection:
            sessions = connection.execute("SELECT * FROM sessions ORDER BY created_ms, id").fetchall()
            for row in sessions:
                if row["kind"] == "practice" and not filters.include_practice:
                    continue
                if filters.participant_code and row["participant_code"] != filters.participant_code:
                    continue
                session = self.experiment.session_snapshot(connection, dict(row))
                if session["kind"] == "experiment" and (not filters.completed_only or session["complete"]):
                    terminal = [trial for trial in session["trials"] if trial["status"] in ("correct", "incomplete")]
                    registry.append({
                        "participant_code": session["participant_code"],
                        "session_id": session["id"],
                        "sequence_no": session["sequence_no"],
                        "created_ms": session["created_ms"],
                        "complete": session["complete"],
                        "completion_code": session.get("completion_code"),
                        "trials": len(session["trials"]),
                        "correct": sum(trial["status"] == "correct" for trial in terminal),
                        "incomplete": sum(trial["status"] == "incomplete" for trial in terminal),
                    })
                selected = {}
                trials_by_id = {trial["id"]: trial for trial in session["trials"]}
                duration_metadata = event_duration_metadata(session["event_log"], trials_by_id)
                base = {"session_id": session["id"], "participant_code": session["participant_code"], "kind": session["kind"]}
                for trial in session["trials"]:
                    if filters.mode and trial["mode"] != filters.mode:
                        continue
                    if filters.level and trial["level"] != filters.level:
                        continue
                    if filters.completed_only and trial["status"] not in ("correct", "incomplete"):
                        continue
                    common = {**base, "trial_id": trial["id"], "mode": trial["mode"], "task_id": trial["task_id"], "level": trial["level"]}
                    selected[trial["id"]] = common
                    actual, analysis = trial["metrics"]["actual"], trial["metrics"]["analysis"]
                    trials.append({**common, "sequence_no": session["sequence_no"], "position": trial["position"], "block_index": trial["block_index"],
                                   "tci": trial["tci"], "status": trial["status"], "end_reason": trial["end_reason"], "gave_up": trial["gave_up"],
                                   "trial_limit_seconds": trial["trial_limit_seconds"], "attempt_limit": trial["attempt_limit"], "wording_version": trial['wording_version'],
                                   "started_ms": trial["started_ms"], "ended_ms": trial["ended_ms"],
                                   "elapsed_ms": actual["elapsed_ms"], "Tcorrect_actual_ms": actual["Tcorrect_ms"], "Tfirst_ms": actual["Tfirst_ms"],
                                   "Tuser_active_ms": actual["Tuser_active_ms"], "A1_actual": actual["A1"], "attempts": actual["attempts"], "Nretry_actual": actual["Nretry"],
                                   "Tcorrect_analysis_ms": analysis["Tcorrect_ms"], "A1_analysis": analysis["A1"], "Nretry_analysis": analysis["Nretry"],
                                   "incomplete": trial["metrics"]["incomplete"]})
                for attempt in session["attempt_log"]:
                    if attempt["trial_id"] not in selected:
                        continue
                    attempts.append({**selected[attempt["trial_id"]], "attempt_id": attempt["id"], "ordinal": attempt["ordinal"], "correct": attempt["correct"],
                                     "started_ms": attempt["started_ms"], "finished_ms": attempt["finished_ms"], "Texec_ms": attempt["Texec_ms"],
                                     "query": attempt["query"], "record_ids": attempt["result"]["record_ids"], "aggregate": attempt["result"]["value"],
                                     "selected_groups": attempt["result"]["selected_groups"], "export_url": attempt["export_url"]})
                for event in session["event_log"]:
                    if event["trial_id"] not in selected:
                        continue
                    payload = event["payload"]
                    trial = trials_by_id[event["trial_id"]]
                    duration = duration_metadata[(event["trial_id"], event["event_id"])]
                    events.append({**selected[event["trial_id"]], "trial_position": trial["position"],
                                   "event_id": event["event_id"], "sequence": event["sequence"], "offset_ms": event["offset_ms"],
                                   "occurred_ms": trial["started_ms"] + event["offset_ms"], **duration,
                                   "received_ms": event["received_ms"], "event_kind": payload["kind"], "target": payload["target"],
                                   "action": payload.get("action"), "x": payload.get("x"), "y": payload.get("y"), "key": payload.get("key"),
                                   "request_id": payload["request_id"]})
                for item in session["interpretation_log"]:
                    if item["trial_id"] not in selected:
                        continue
                    interpretations.append({**selected[item["trial_id"]], "request_id": item["request_id"], "user_text": item["user_text"],
                                            "status": item["status"], "requested_model": item["requested_model"],
                                            "response_model": item["response_model"], "provider": item["provider"],
                                            "system_fingerprint": item["system_fingerprint"], "prompt_version": item["prompt_version"],
                                            "prompt_sha256": item["prompt_sha256"], "temperature": item["temperature"],
                                            "raw_response": item["raw_response"], "query": item["query"], "error_code": item["error_code"],
                                            "error_message": item["error_message"], "Tllm_ms": item["llm_ms"],
                                            "input_tokens": item["input_tokens"], "output_tokens": item["output_tokens"],
                                            "started_ms": item["started_ms"], "finished_ms": item["finished_ms"]})
                for item in session["m4_transcription_log"]:
                    if item["trial_id"] not in selected:
                        continue
                    m4_transcriptions.append({**selected[item["trial_id"]], "request_id": item["request_id"],
                                              "mime_type": item["mime_type"], "audio_bytes": item["audio_bytes"],
                                              "audio_duration_ms": item["audio_duration_ms"], "requested_model": item["requested_model"],
                                              "compute_type": item["compute_type"], "requested_language": item["requested_language"],
                                              "detected_language": item["detected_language"], "language_probability": item["language_probability"],
                                              "transcript": item["transcript"], "Tasr_ms": item["asr_ms"],
                                              "started_ms": item["started_ms"], "finished_ms": item["finished_ms"]})
                for item in session["agent_log"]:
                    if item["trial_id"] not in selected:
                        continue
                    agent_runs.append({**selected[item["trial_id"]], "request_id": item["request_id"], "user_text": item["user_text"],
                                       "status": item["status"], "requested_model": item["requested_model"],
                                       "response_model": item["response_model"], "provider": item["provider"],
                                       "prompt_version": item["prompt_version"], "prompt_sha256": item["prompt_sha256"],
                                       "temperature": item["temperature"], "max_steps": item["max_steps"], "query": item["query"],
                                       "trajectory": item["trajectory"], "final_text": item["final_text"], "termination": item["termination"],
                                       "error_code": item["error_code"], "error_message": item["error_message"],
                                       "Tllm_ms": item["llm_ms"], "Ttool_ms": item["tool_ms"], "Tagent_ms": item["total_ms"],
                                       "llm_calls": item["llm_calls"], "tool_calls": item["tool_calls"],
                                       "input_tokens": item["input_tokens"], "output_tokens": item["output_tokens"],
                                       "started_ms": item["started_ms"], "finished_ms": item["finished_ms"]})
            generated = self.experiment.clock()
        return {"protocol": {"export_schema_version": 8, "generated_at_ms": generated, "filters": filters.model_dump(mode="json"),
                             "manifest": self.experiment.manifest, "time_units": "milliseconds", "missing_value": "empty cell / JSON null",
                             "event_duration_semantics": "point/closing events=0; interval-start events span to matching finish/close; unclosed spans are right-censored at the observation boundary",
                             "event_ordering": "filter by participant_code; order by occurred_ms, trial_position, sequence",
                             "summary_population": "terminal trials only; actual Tcorrect includes successful trials only",
                             "training_excluded": not filters.include_practice},
                "registry": registry, "trials": trials, "attempts": attempts, "events": events, "interpretations": interpretations,
                "m4_transcriptions": m4_transcriptions, "agent_runs": agent_runs, "summary": summarize(trials)}


def _mean(rows, field):
    values = [row[field] for row in rows if row[field] is not None]
    return mean(values) if values else None


def summarize(trials: list[dict]) -> list[dict]:
    grouped = defaultdict(list)
    for trial in trials:
        grouped[(trial["kind"], trial["mode"], trial["level"], trial['trial_limit_seconds'], trial['attempt_limit'], trial['wording_version'])].append(trial)
    result = []
    for (kind, mode, level, time_limit, attempt_limit, wording), rows in sorted(grouped.items()):
        terminal = [row for row in rows if row["status"] in ("correct", "incomplete")]
        success = [row for row in terminal if row["status"] == "correct"]
        result.append({"kind": kind, "mode": mode, "level": level, "participants": len({row["participant_code"] for row in terminal}),
                       'trial_limit_seconds': time_limit, 'attempt_limit': attempt_limit, 'wording_version': wording,
                       "trials": len(rows), "pending": sum(row["status"] == "pending" for row in rows), "active": sum(row["status"] == "active" for row in rows),
                       "completed": len(terminal), "correct": len(success), "incomplete": len(terminal) - len(success),
                       "success_rate": len(success) / len(terminal) if terminal else None,
                       "A1_observed_n": sum(row["A1_actual"] is not None for row in terminal),
                       "A1_actual_rate": _mean(terminal, "A1_actual"), "A1_analysis_rate": _mean(terminal, "A1_analysis"),
                       "Tcorrect_actual_n": len(success), "Tcorrect_actual_mean_ms": _mean(success, "Tcorrect_actual_ms"),
                       "Tcorrect_analysis_mean_ms": _mean(terminal, "Tcorrect_analysis_ms"),
                       "Tuser_active_mean_ms": _mean(terminal, "Tuser_active_ms"),
                       "Nretry_actual_mean": _mean(terminal, "Nretry_actual"), "Nretry_analysis_mean": _mean(terminal, "Nretry_analysis")})
    return result


def cell(value):
    if isinstance(value, (dict, list)):
        value = encode(value)
    # Spreadsheet programs must treat participant/browser data as text.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def table_csv(snapshot: dict, name: str) -> str:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    columns = TABLE_COLUMNS[name]
    writer.writerow(columns)
    for row in snapshot[name]:
        writer.writerow([cell(row.get(column)) for column in columns])
    return "\ufeff" + stream.getvalue()


def csv_bundle(snapshot: dict) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for table in TABLE_COLUMNS:
            archive.writestr(f"{table}.csv", table_csv(snapshot, table).encode("utf-8"))
        archive.writestr("protocol.json", encode(snapshot["protocol"]).encode("utf-8"))
    return buffer.getvalue()


def excel_workbook(snapshot: dict) -> bytes:
    workbook = Workbook(write_only=True)
    protocol = workbook.create_sheet("protocol")
    protocol.append(["parameter", "value"])
    for key, value in snapshot["protocol"].items():
        protocol.append([key, excel_cell(value)])
    for table, columns in TABLE_COLUMNS.items():
        sheet = workbook.create_sheet(table)
        sheet.append(columns)
        for index, row in enumerate(snapshot[table]):
            # Stay below Excel's row limit, including the header.
            if index and index % 1_000_000 == 0:
                sheet = workbook.create_sheet(f"{table}_{index // 1_000_000 + 1}")
                sheet.append(columns)
            sheet.append([excel_cell(row.get(column)) for column in columns])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def excel_cell(value):
    value = cell(value)
    if isinstance(value, str):
        # Preserve control characters explicitly instead of losing them in XLSX.
        value = "".join(f"\\u{ord(ch):04x}" if ord(ch) < 32 and ch not in "\t\n\r" else ch for ch in value)
        if len(value) > 32767:
            raise DomainError("excel_cell_limit", "Значение превышает лимит ячейки Excel. Используйте CSV/ZIP без потери данных.")
    return value
